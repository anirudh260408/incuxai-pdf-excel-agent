"""
INCUXAI AI AGENT: Comprehensive Accuracy Benchmark & Validation Suite.
Evaluates:
1. Multi-Page Continuity Benchmark (4-Page Financial Statement)
2. 10+ Column Alignment & Arithmetic Benchmark (3-Page 12-Column Supply Chain Report)
3. Scanned OCR Reconstruction Benchmark (Scanned Logistics Invoice)
4. 50-Page Stress Test (Validates handling up to 50 pages as specified in Requirement 1)
"""

import os
import time
import openpyxl
import pandas as pd
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, PageBreak
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib import colors

from agent.pipeline import DocumentPipeline


def generate_50_page_stress_pdf(output_path: str, rows_per_page: int = 20) -> str:
    """Generate a realistic 50-page continuous tabular report for stress testing."""
    print("Generating 50-page PDF stress document...")
    doc = SimpleDocTemplate(
        output_path,
        pagesize=letter,
        leftMargin=36,
        rightMargin=36,
        topMargin=36,
        bottomMargin=36
    )
    styles = getSampleStyleSheet()
    elements = []

    headers = ["Record ID", "Client Name", "Region", "Transaction Code", "Amount ($)", "Status"]
    statuses = ["Completed", "Pending", "Audited", "Settled"]

    for page_num in range(1, 51):
        page_rows = [headers]
        for r in range(1, rows_per_page + 1):
            rec_id = f"REC-{(page_num - 1) * rows_per_page + r:05d}"
            client = f"Enterprise Client {(r * page_num) % 97 + 1}"
            region = ["North America", "EMEA", "APAC", "LATAM"][r % 4]
            txn_code = f"TX-{(page_num * 100 + r)}"
            amt = f"{(r * 123.45):,.2f}"
            status = statuses[r % len(statuses)]
            page_rows.append([rec_id, client, region, txn_code, amt, status])

        t = Table(page_rows, colWidths=[75, 140, 90, 80, 80, 75])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#1B365D')),
            ('TEXTCOLOR', (0, 0), (-1, 0), colors.white),
            ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, 0), 8),
            ('GRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#D3D9DE')),
            ('ROWBACKGROUNDS', (0, 1), (-1, -1), [colors.white, colors.HexColor('#F7F9FC')])
        ]))
        elements.append(Paragraph(f"<b>ANNUAL AUDIT LOG - 50 PAGE MASTER RUN</b> (Page {page_num} of 50)", styles['Normal']))
        elements.append(t)
        if page_num < 50:
            elements.append(PageBreak())

    doc.build(elements)
    print(f"50-page PDF successfully created: {output_path}")
    return output_path


def run_full_evaluation():
    print("=" * 80)
    print("🔬 RUNNING COMPREHENSIVE INCUXAI BENCHMARK & ACCURACY SUITE")
    print("=" * 80)

    pipeline = DocumentPipeline()
    samples_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "samples")
    out_dir = os.path.join(samples_dir, "converted")
    os.makedirs(out_dir, exist_ok=True)

    test_cases = [
        {
            "name": "Sample 1: Multi-Page Financial Statement",
            "path": os.path.join(samples_dir, "sample_1_financial_statement.pdf"),
            "expected_pages": 4,
            "expected_cols": 7,
            "min_expected_rows": 55,
            "category": "Multi-page Continuity & Repeated Header Removal"
        },
        {
            "name": "Sample 2: Master Inventory & Order Report",
            "path": os.path.join(samples_dir, "sample_2_inventory_order_report.pdf"),
            "expected_pages": 3,
            "expected_cols": 12,
            "min_expected_rows": 36,
            "category": "10+ Columns & Arithmetic Error Detection"
        },
        {
            "name": "Sample 3: Scanned Logistics Tax Invoice",
            "path": os.path.join(samples_dir, "sample_3_scanned_invoice.pdf"),
            "expected_pages": 1,
            "expected_cols": 6,
            "min_expected_rows": 6,
            "category": "OCR Engine & Geometric Table Reconstruction"
        }
    ]

    benchmark_results = []

    for tc in test_cases:
        print(f"\nEvaluating: {tc['name']}...")
        start = time.time()
        base_name = os.path.splitext(os.path.basename(tc['path']))[0]
        out_excel = os.path.join(out_dir, f"{base_name}.xlsx")
        res = pipeline.process_document(tc['path'], output_path=out_excel)
        duration = round(time.time() - start, 3)

        # Check excel output
        wb = openpyxl.load_workbook(out_excel)
        has_summary_sheet = "Audit & Validation Summary" in wb.sheetnames
        data_sheet = wb["Table 1"] if "Table 1" in wb.sheetnames else wb.worksheets[1]

        actual_cols = data_sheet.max_column
        actual_rows = data_sheet.max_row - 1  # minus header

        # Metrics
        col_accuracy = 100.0 if actual_cols == tc["expected_cols"] else max(0.0, 100 - abs(actual_cols - tc["expected_cols"]) * 10)
        row_accuracy = min(100.0, (actual_rows / tc["min_expected_rows"]) * 100.0)
        page_continuity = 100.0 if res.total_pages == tc["expected_pages"] and len(res.tables_data) == 1 else 90.0

        benchmark_results.append({
            "Test Case": tc["name"],
            "Category": tc["category"],
            "Pages": res.total_pages,
            "Cols (Found/Exp)": f"{actual_cols}/{tc['expected_cols']}",
            "Rows (Found/Exp)": f"{actual_rows}/{tc['min_expected_rows']}",
            "Integrity Score": f"{res.overall_integrity_score}%",
            "Audit Findings": f"{res.total_errors} err, {res.total_warnings} warn",
            "Speed": f"{duration}s",
            "Verdict": "PASS" if res.success and col_accuracy >= 95 and row_accuracy >= 95 else "FAIL"
        })

    # Run 50-Page Stress Test
    print("\n" + "-" * 80)
    print("🚀 RUNNING 50-PAGE LARGE DOCUMENT STRESS TEST (Challenge Requirement 1)")
    print("-" * 80)
    stress_pdf = os.path.join(samples_dir, "stress_test_50_pages.pdf")
    stress_excel = os.path.join(out_dir, "stress_test_50_pages.xlsx")

    generate_50_page_stress_pdf(stress_pdf, rows_per_page=20)

    t0 = time.time()
    stress_res = pipeline.process_document(stress_pdf, output_path=stress_excel)
    stress_duration = round(time.time() - t0, 3)

    stress_wb = openpyxl.load_workbook(stress_excel)
    stress_sheet = stress_wb["Table 1"] if "Table 1" in stress_wb.sheetnames else stress_wb.worksheets[1]
    stress_cols = stress_sheet.max_column
    stress_rows = stress_sheet.max_row - 1

    benchmark_results.append({
        "Test Case": "50-Page Enterprise Audit Log Stress Test",
        "Category": "Scalability & 50-Page Table Stitching",
        "Pages": stress_res.total_pages,
        "Cols (Found/Exp)": f"{stress_cols}/6",
        "Rows (Found/Exp)": f"{stress_rows}/1000",
        "Integrity Score": f"{stress_res.overall_integrity_score}%",
        "Audit Findings": f"{stress_res.total_errors} err, {stress_res.total_warnings} warn",
        "Speed": f"{stress_duration}s ({round(stress_duration/50, 3)}s/page)",
        "Verdict": "PASS" if stress_res.success and stress_rows >= 1000 else "FAIL"
    })

    print("\n" + "=" * 80)
    print("🏆 FINAL BENCHMARK & ACCURACY SCORECARD")
    print("=" * 80)
    df_results = pd.DataFrame(benchmark_results)
    print(df_results.to_string(index=False))
    print("=" * 80)

    return df_results


if __name__ == "__main__":
    run_full_evaluation()
